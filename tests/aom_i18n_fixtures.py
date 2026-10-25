"""Deterministické (anonymizované) vstupy pre zlaté testy posudku — tvar podľa reálneho
econ_results.full_response z prodú (2026-10-25), čísla syntetické."""
import datetime as dt
import io
import math
import re

_PV_SHAPE = [0.031, 0.049, 0.083, 0.109, 0.125, 0.132, 0.139, 0.125, 0.090, 0.058, 0.033, 0.026]
_LOAD_SHAPE = [0.082, 0.090, 0.093, 0.076, 0.069, 0.077, 0.085, 0.063, 0.074, 0.104, 0.103, 0.084]


def _variant(vid, label, pv_kwp, bess_kwh, capex_pv, capex_bess, saving, npv, irr, pb,
             direct, battery, exported, curtailed, merchant=None, load_mwh=388.0):
    pv_kwh = pv_kwp * 1035.0
    capex = capex_pv + capex_bess
    cf = [-capex] + [round(saving * (0.995 ** y), 2) for y in range(20)]
    monthly = [{"month": m + 1, "pv_kwh": pv_kwh * _PV_SHAPE[m], "load_kwh": load_mwh * 1000 * _LOAD_SHAPE[m],
                "total_eur": saving * _PV_SHAPE[m], "export_kwh": pv_kwh * _PV_SHAPE[m] * exported / 100,
                "import_kwh": load_mwh * 1000 * _LOAD_SHAPE[m] * 0.75, "arbitrage_eur": 0, "peak_shaving_eur": 0,
                "solar_export_eur": 60 * pv_kwh * _PV_SHAPE[m] * exported / 100000,
                "solar_to_load_eur": saving * _PV_SHAPE[m] * 0.95} for m in range(12)]
    vs = {"total_eur": saving, "merchant_eur": (merchant or {}).get("merchant_net_eur", 0), "arbitrage_eur": 0,
          "bess_total_eur": 900.0 if bess_kwh else 0, "peak_shaving_eur": 0, "site_savings_eur": saving,
          "solar_export_eur": saving * 0.04, "merchant_revenue_eur": (merchant or {}).get("merchant_net_eur", 0),
          "mrk_penalty_avoided_eur": 0, "total_project_value_eur": saving,
          "bess_self_consumption_eur": 1100.0 if bess_kwh else 0, "solar_self_consumption_eur": saving * 0.9}
    ef = {"pv_total_mwh": pv_kwh / 1000, "pv_to_bat_mwh": pv_kwh * battery / 100000,
          "load_total_mwh": load_mwh, "pv_to_grid_mwh": pv_kwh * exported / 100000,
          "pv_to_load_mwh": pv_kwh * direct / 100000, "bat_to_load_mwh": pv_kwh * battery * 0.9 / 100000,
          "grid_export_mwh": pv_kwh * exported / 100000, "grid_import_mwh": load_mwh * 0.75,
          "grid_to_bat_mwh": 12.0 if bess_kwh else 0, "grid_to_load_mwh": load_mwh * 0.66,
          "pv_curtailed_mwh": pv_kwh * curtailed / 100000}
    return {"variant_id": vid, "db_id": vid, "label": label, "pv_kwp": pv_kwp, "bess_kwh": bess_kwh,
            "bess_kw": bess_kwh / 2, "inverter_kw": round(pv_kwp * 0.95, 1), "irr_pct": irr, "npv_eur": npv,
            "dotacia_eur": 0, "capex_pv_eur": capex_pv, "capex_bess_eur": capex_bess, "capex_total_eur": capex,
            "net_capex_eur": capex, "capex_source": "manual", "ems_strategy": "rule_based",
            "lcoe_eur_mwh": 61.2, "lcos_eur_mwh": 350.8 if bess_kwh else None, "pv_total_kwh": pv_kwh,
            "saving_y1_eur": saving, "site_savings_eur": saving, "value_streams": vs, "cashflow_array": cf,
            "grid_import_kwh": load_mwh * 750, "merchant_detail": merchant or {}, "monthly_summary": monthly,
            "payback_simple_y": pb, "samospotreba_pct": direct + battery, "samostatnost_pct": 25.7,
            "merchant_revenue_eur": vs["merchant_revenue_eur"], "total_project_value_eur": saving,
            "solar_consumption_pct": {"exported": exported, "curtailed": curtailed, "direct_to_load": direct,
                                      "charging_battery": battery},
            "energy_flow": ef, "carbon": {"trees_equivalent": 1327, "barrels_oil_avoided": 64,
                                          "co2_avoided_t_per_year": 27.9},
            "rank_labels": ["Najvyššie NPV", "Najvyššie IRR", "Najrýchlejšia návratnosť"]}


_TARIFF = {"rok": 2026, "sadzba": "VN", "typ_tarify": "fix", "njf_eur_mwh": 0, "tps_eur_mwh": 14.07,
           "tss_eur_mwh": 7.79, "distribuutor": "SSE", "silova_eur_mwh": 134.87, "straty_eur_mwh": 5.2,
           "distrib_eur_mwh": 16.55, "spot_avg_eur_mwh": 102.86, "obchodnik_eur_mwh": 8,
           "regulovane_eur_mwh": 51.39, "export_price_eur_mwh": 60, "retail_total_eur_mwh": 194.26,
           "spotrebna_dan_eur_mwh": 7.79, "rk_kapacita_eur_kw_mes": 0, "mrk_kapacita_eur_kw_mes": 6.764,
           "mrk_export_penalty_eur_mwh": 12.5}

_MERCHANT = {"throughput_mwh": 310.5, "equiv_cycles": 276.0, "buy_eur": 21000.0, "sell_eur": 52000.0,
             "gross_eur": 31000.0, "organizer_fee_eur": 4650.0, "fee_pct": 15.0, "merchant_net_eur": 26350.0}


def _analyza(variants, sel, **kw):
    a = {"id": "00000000-0000-4000-8000-000000000001", "name": "Testovací závod — OM Testovo",
         "customers": {"company_name": "Testovacia firma s.r.o.", "first_name": None, "last_name": None,
                       "email": None, "ico": "12345678"},
         "om_psc": "97243", "om_address": "Priemyselná 1, 972 43 Testovo", "om_rk_kw": 210, "om_mrk_kw": 240,
         "om_sadzba": "VN", "om_tarif_typ": "fix", "om_distributor": "SSE", "tarif_buy": 0.1811, "tarif_sell": 0.06,
         "tarif_silova_eur_mwh": 134.87, "tarif_distribucia_eur_mwh": 16.554, "tarif_tps_eur_mwh": 14.071,
         "tarif_oze_eur_mwh": 0, "tarif_ostatne_eur_mwh": 15.57, "tarif_mrk_kapacita_eur_mw_mes": 6764,
         "scenario_type": "fve_bess", "dotacia_scheme": "zelena_podnikom", "dotacia_enabled": False,
         "merchant_mode": False, "merchant_organizer_fee_pct": 15, "consumption_avg_kw": 44.3,
         "consumption_peak_kw_15min": 226, "consumption_peak_kw_hourly": 219.7, "consumption_annual_mwh": 388.06,
         "consumption_strategy": "measured", "fve_tilt_deg": 13, "fve_azimuth_deg": 180, "fve_topology": "south",
         "ma_danovy_zaklad": True, "savings_coefficient": 1, "existing_fve_kwp": 0, "existing_bess_kwh": 0,
         "selected_variant_id": sel, "workspace": "b2b",
         "econ_results": {"full_response": {"variants": variants, "tariff_detail": _TARIFF}}}
    a.update(kw)
    return a


def _site_variants():
    return [
        _variant("V1", "FVE 119,84 kWp + BESS 112,5 kWh", 119.84, 112.5, 71476.0, 54000.0, 19185.8, 70663.2,
                 13.14, 6.16, 75.9, 5.2, 9.6, 9.3),
        _variant("V2", "FVE 100 kWp", 100.0, 0, 99280.0, 0, 16366.9, 71851.3, 14.92, 5.65, 77.4, 0, 12.7, 9.9),
        _variant("V3", "FVE 150 kWp", 150.0, 0, 148920.0, 0, 23678.2, 99864.6, 14.33, 5.81, 73.7, 0, 15.8, 10.5),
    ]


def _hourly():
    """Syntetický hodinový profil 8760 h: (hodina, víkend, kW, mesiac) ako render_posudok_chocosuc."""
    out, d0 = [], dt.datetime(2025, 1, 1)
    for h in range(8760):
        d = d0 + dt.timedelta(hours=h)
        kw = 44 + 30 * math.sin((d.hour - 6) / 24 * 2 * math.pi) + (0 if d.weekday() < 5 else -15)
        out.append((d.hour, d.weekday() >= 5, round(max(kw, 5.0), 2), d.month))
    return out


def fixture_analyza():
    return _analyza(_site_variants(), "V1")


def fixture_analyza_merchant():
    vs = _site_variants()
    m = _variant("M1", "FVE 725 kWp + BESS 1 MWh (bilančná skupina)", 725.46, 1000.0, 420000.0, 318000.0,
                 112000.0, 610000.0, 15.2, 6.4, 62.0, 4.0, 30.0, 4.0, merchant=_MERCHANT, load_mwh=2400.0)
    return _analyza([m] + vs[1:], "M1", merchant_mode=True, om_tarif_typ="spot", consumption_annual_mwh=2400.0,
                    consumption_avg_kw=274.0, consumption_peak_kw_15min=980, om_mrk_kw=1200, om_rk_kw=1000)


def _ctx(a):
    from posudok_chocosuc.context import build_chocosuc_context
    return build_chocosuc_context(a, a["econ_results"]["full_response"]["variants"], hourly=_hourly(), selected_row=None)


def fixture_ctx():
    return _ctx(fixture_analyza())


def fixture_ctx_bess_merchant():
    return _ctx(fixture_analyza_merchant())


_IMG = re.compile(r'data:image/(?:png|jpeg|svg\+xml);base64,[A-Za-z0-9+/=]+')
_DATE = re.compile(r'\b\d{2}\.\d{2}\.\d{4}\b')
_SCRIPT = re.compile(r"<script>.*?</script>", re.S)


def normalize_html(s: str, strip_scripts: bool = False) -> str:
    s = _IMG.sub("data:IMG", s)
    if strip_scripts:
        s = _SCRIPT.sub("<script>…</script>", s)
    return _DATE.sub("DD.MM.YYYY", s)


def docx_text(b: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(b))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            parts.append(" | ".join(c.text for c in r.cells))
    return _DATE.sub("DD.MM.YYYY", "\n".join(parts))


_CD = re.compile(r"var CD=(\{.*?\});Chart\.defaults", re.S)


def _chart_texts(fn):
    """Zachytí texty matplotlib grafov (nadpisy, osi, legendy, ticky) namiesto PNG."""
    import matplotlib.text
    from posudok_chocosuc import charts as C
    seen, orig = [], C._b64

    def rec(fig):
        fig.canvas.draw()
        seen.append([t.get_text() for t in fig.findobj(matplotlib.text.Text) if t.get_text().strip()])
        return orig(fig)
    C._b64 = rec
    try:
        out = fn()
    finally:
        C._b64 = orig
    return out, seen


def render_outputs():
    """Všetky porovnávané výstupy posudku v aktuálnom jazyku (aom_i18n.use_lang okolo volania)."""
    import json
    from posudok_chocosuc.generator import render_chocosuc_html
    from posudok_chocosuc.generator_docx import generate_chocosuc_docx
    out = {}
    for tag, f in (("site", fixture_ctx), ("merchant", fixture_ctx_bess_merchant)):
        html, texts = _chart_texts(lambda: render_chocosuc_html(f()))
        m = _CD.search(html)
        out[f"posudok_{tag}.html"] = normalize_html(html, strip_scripts=True).replace('<html lang="sk">', "<html>")
        out[f"posudok_{tag}_chartjs.json"] = json.dumps(json.loads(m.group(1)) if m else None, ensure_ascii=False, indent=1)
        out[f"posudok_{tag}_charts.json"] = json.dumps(texts, ensure_ascii=False, indent=1)
        out[f"posudok_{tag}_docx.txt"] = docx_text(generate_chocosuc_docx(f()))
    return out


_SK_CHARS = re.compile(r"[A-Za-zÀ-ž]*[áčďéíľĺňóôŕšťúýžÁČĎÉÍĽĹŇÓÔŔŠŤÚÝŽ][A-Za-zÀ-ž]*")
_DATA_WORDS = {"Priemyselná", "Zelená", "ÚRSO", "OKTE"}  # dáta fixtúry a glosár (pôvodné SK názvy)


def slovak_leftovers(text: str) -> list:
    """Slová so slovenskou diakritikou vo viditeľnom texte (bez <script>/<style>/značiek) — pre EN/DE výstup."""
    t = re.sub(r"<(script|style)\b.*?</\1>", " ", text, flags=re.S)
    t = re.sub(r"<[^>]+>", " ", t)
    return sorted({w for w in _SK_CHARS.findall(t) if w not in _DATA_WORDS})
