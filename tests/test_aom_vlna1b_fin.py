"""Regresné testy — audit AOM 2026-10-24, vlna 1b: finančné nálezy A14, A19, A20, A21
(docs/aom-audit-2026-10-24/gpt_findings.md, 05-financie/report.md).

Každý test pred opravou padal (reprodukcia nálezu), po oprave prechádza.
Spúšťanie: `python3 tests/test_aom_vlna1b_fin.py` alebo `pytest tests/test_aom_vlna1b_fin.py`.
"""
import io
import os
import sys
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from test_aom_vlna1 import ROOT, _req  # noqa: E402  (nastaví sys.path + ENERGO_* env)

_CACHE: dict = {}


def _run(pv, bess, **kw):
    """Plný beh pipeline → (VariantResult, API dict variantu, celá API odpoveď). Cachované."""
    key = (pv, bess, tuple(sorted(kw.items())))
    if key not in _CACHE:
        from energovision_analytics.api.services.engine_service import (
            run_variants_pipeline, build_run_variants_response)
        rq = _req(**kw)
        rq["variants"]["pv_kwp_options"] = [pv]
        rq["variants"]["bess_kwh_options"] = [bess]
        out = run_variants_pipeline(rq)
        resp = build_run_variants_response(out)
        _CACHE[key] = (out["results"][0], resp["variants"][0], resp)
    return _CACHE[key]


# ---------- A14: koeficient úspory aj v prezentovanej úspore / value streams / mesiacoch ----------
def _a14_check(pv, bess):
    r1, v1, _ = _run(pv, bess)
    r2, v2, _ = _run(pv, bess, savings_coefficient=1.25)
    y1_rev = r2.financial.yearly_cashflows[1].revenue_total
    # prezentovaná ročná úspora = príjem roku 1 v cashflow (pred zmenou: 35 025 vs 43 781 €)
    assert abs(v2["saving_y1_eur"] - y1_rev) < 1.0, (v2["saving_y1_eur"], y1_rev)
    assert abs(v2["saving_y1_eur"] - 1.25 * v1["saving_y1_eur"]) < 1.0, (v1["saving_y1_eur"], v2["saving_y1_eur"])
    # value streams aj ich súčet na rovnakej báze ako cashflow
    assert abs(v2["value_streams"]["total_eur"] - y1_rev) < 1.0, (v2["value_streams"]["total_eur"], y1_rev)
    for k, a in v1["value_streams"].items():
        assert abs(v2["value_streams"][k] - 1.25 * a) < 0.01, (k, a, v2["value_streams"][k])
    # mesačný súhrn škáluje rovnakým koeficientom (súčet mesiacov rieši A13 samostatne)
    for m1, m2 in zip(v1["monthly_summary"], v2["monthly_summary"]):
        assert abs(m2["total_eur"] - 1.25 * m1["total_eur"]) < 0.01, (m1, m2)
    # raw hodnota ostáva dostupná na audit
    assert abs(r2.financial.annual_saving_y1_raw_eur - v1["saving_y1_eur"]) < 1.0
    # koef 1.0 → bez zmeny (raw == prezentovaná)
    assert abs(r1.financial.annual_saving_y1_eur - r1.financial.annual_saving_y1_raw_eur) < 1e-9


def test_a14_koeficient_fve_bess():
    _a14_check(300, 200)


def test_a14_koeficient_pv_only():
    # PV-only ide cez PVGIS vetvu mesačného súhrnu (bez EMS monthly_flows)
    _a14_check(300, 0)


def test_a14_cashflow_builder_priamo():
    from energovision_analytics.financial import CashflowBuilder
    b = CashflowBuilder(capex_solar_eur=100000, capex_bess_eur=0, savings_coefficient=0.8)
    f = b.build(annual_saving_y1_eur=17102.48, saving_decomp_y1={"sav_solar_self_cons_eur": 17102.48})
    assert abs(f.annual_saving_y1_eur - f.yearly_cashflows[1].revenue_total) < 1e-6, (
        f.annual_saving_y1_eur, f.yearly_cashflows[1].revenue_total)
    assert abs(f.annual_saving_y1_eur - 13681.984) < 0.01
    assert f.annual_saving_y1_raw_eur == 17102.48


# ---------- A19: variant bez investície neplatí monitoring ani OPEX ----------
def test_a19_nulovy_variant_bez_opex():
    from energovision_analytics.financial import CashflowBuilder
    f = CashflowBuilder(capex_solar_eur=0, capex_bess_eur=0).build(
        annual_saving_y1_eur=0.0, saving_decomp_y1={})
    assert all(c.opex_total == 0 for c in f.yearly_cashflows), [c.opex_total for c in f.yearly_cashflows]
    assert f.npv_eur == 0.0, f.npv_eur  # pred opravou −3 440,98 € (300 €/r monitoring)


def test_a19_engine_0_kwp_0_kwh():
    _, v, _ = _run(0, 0)
    assert abs(v["npv_eur"]) < 0.01, v["npv_eur"]
    assert all(abs(x) < 0.01 for x in v["cashflow_array"]), v["cashflow_array"]


def test_a19_s_technologiou_monitoring_ostava():
    from energovision_analytics.financial import CashflowBuilder
    f = CashflowBuilder(capex_solar_eur=100000, capex_bess_eur=0).build(
        annual_saving_y1_eur=0.0, saving_decomp_y1={})
    assert f.yearly_cashflows[1].cost_monitoring == 300


# ---------- A20: presne 80 % samospotreby → 45 % intenzita ----------
def test_a20_dotacia_presne_80():
    from energovision_analytics.financial.dotacie import apply_dotacia, load_dotacie_schemes
    sch = load_dotacie_schemes()
    res = apply_dotacia(scheme_id="zelena_podnikom", capex_eur=100000, samospotreba_pct=80.0,
                        project_type="FVE", schemes=sch, installed_kw=100)
    assert res["eligible"] and abs(res["amount_eur"] - 45000) < 0.01, res["amount_eur"]  # pred: 40 000
    res = apply_dotacia(scheme_id="zelena_podnikom", capex_eur=100000, samospotreba_pct=79.99,
                        project_type="FVE", schemes=sch, installed_kw=100)
    assert abs(res["amount_eur"] - 40000) < 0.01, res["amount_eur"]


def test_a20_legacy_helper_presne_80():
    from energovision_analytics.financial import sk_dotacia_zelena_podnikom
    assert abs(sk_dotacia_zelena_podnikom(100000, 80.0)["amount_eur"] - 45000) < 0.01


# ---------- A21: jeden zdroj ekonomických defaultov, DPPO 21 % všade ----------
def test_a21_yaml_sa_nacita():
    from energovision_analytics.core import defaults
    p = Path(defaults._DEFAULT_YAML)
    assert p.exists(), p  # pred: repos/data/config/... neexistuje → zabudované defaulty
    assert p == Path(ROOT) / "aom_data" / "config" / "economic_defaults.yaml"
    e = defaults.get_econ_defaults(p)
    assert e._source_path == str(p)
    assert e.dppo.default_pct == 0.21 and e.dppo.standard_pct == 0.21, e.dppo
    # YAML a zabudované defaulty sa nerozchádzajú (fallback dáva rovnaké čísla)
    assert e == defaults.EconomicDefaults(_source_path=str(p)), e


def test_a21_manifest_hash():
    from energovision_analytics.core.run_manifest import build_run_manifest
    m = build_run_manifest(tariff_yaml=os.environ["ENERGO_TARIFF_YAML"], spot_csv=os.environ["ENERGO_SPOT_CSV"])
    assert m.economic_defaults_hash != "missing", m.economic_defaults_path


def test_a21_request_schema_dppo():
    from energovision_analytics.api.models.requests import FinancialConfig
    from energovision_analytics.core.defaults import ECON
    assert FinancialConfig().dppo_pct == ECON.dppo.default_pct == 0.21


def test_a21_premium_posudok_text_21():
    from docx import Document
    from energovision_analytics.reporting.posudok_premium import generate_premium_posudok
    _, _, resp = _run(300, 200)
    b = generate_premium_posudok(client_name="Test s.r.o.", project_id="P-TEST", site_meta={},
                                 run_response=resp)
    d = Document(io.BytesIO(b))
    txt = "\n".join(p.text for p in d.paragraphs)
    txt += "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert "DPPO 22" not in txt, [ln for ln in txt.splitlines() if "DPPO" in ln]
    assert "DPPO 21 %" in txt
    sel = resp["variants"][0]
    assert f"~{sel['net_capex_eur'] * 0.21:,.0f} €" in txt


def test_a21_docx_generator_text():
    src = (Path(ROOT) / "energovision_analytics" / "reporting" / "docx_generator.py").read_text()
    assert "DPPO 22" not in src


if __name__ == "__main__":
    import traceback
    fails = 0
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    for n, f in tests:
        try:
            f()
            print(f"PASS {n}")
        except Exception:
            fails += 1
            print(f"FAIL {n}")
            traceback.print_exc()
    print(f"{len(tests) - fails}/{len(tests)} OK")
    sys.exit(1 if fails else 0)
